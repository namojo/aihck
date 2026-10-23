"""산출물: 엑셀(추적표·위험·배점-목차·커버리지), 워드(목차 초안), HTML 대시보드."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .schemas import CATEGORY_LABELS, ReviewReport

WEB_DIR = Path(__file__).parent / "web"
FILL = {
    "high": PatternFill("solid", fgColor="F8D7DA"),
    "medium": PatternFill("solid", fgColor="FFF3CD"),
    "low": PatternFill("solid", fgColor="E2E3E5"),
    "covered": PatternFill("solid", fgColor="D4EDDA"),
    "partial": PatternFill("solid", fgColor="FFF3CD"),
    "missing": PatternFill("solid", fgColor="F8D7DA"),
    "header": PatternFill("solid", fgColor="1B2A49"),
}
OBLIGATION_KO = {"must": "필수", "should": "권고", "may": "선택"}
SEVERITY_KO = {"high": "높음", "medium": "보통", "low": "낮음"}
STATUS_KO = {"covered": "충족", "partial": "부분", "missing": "누락"}


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: list[int], fills: list | None = None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = FILL["header"]
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, r in enumerate(rows):
        ws.append(r)
        if fills and fills[i]:
            for c in ws[ws.max_row]:
                c.fill = fills[i]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    return ws


def write_xlsx(report: ReviewReport, path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    cov = {it.requirement_id: it for it in (report.coverage.items if report.coverage else [])}
    crit_by_req: dict[str, str] = {}
    if report.toc:
        for node in report.toc.toc:
            for rid in node.requirement_ids:
                crit_by_req.setdefault(rid, f"{node.number} {node.title}")
    rows, fills = [], []
    for r in report.requirements:
        c = cov.get(r.id)
        rows.append([r.id, r.title, f"{r.category} {CATEGORY_LABELS.get(r.category, '')}", OBLIGATION_KO[r.obligation], r.quote, r.page, r.section or "",
                     ", ".join(r.deliverables), "표" if r.source == "table" else "본문", crit_by_req.get(r.id, "(미연결)"),
                     STATUS_KO[c.status] if c else "", (c.evidence_quote or c.gap_note or "") if c else "", ", ".join(r.flags)])
        fills.append(FILL.get(c.status) if c else None)
    _sheet(wb, "요구사항 추적표", ["ID", "요구사항", "분류", "필수", "원문", "쪽", "섹션", "산출물", "출처", "제안서 목차", "커버리지", "근거/보완", "비고"],
           rows, [11, 28, 14, 6, 60, 5, 22, 18, 6, 24, 8, 36, 14], fills)

    rows, fills = [], []
    for f in report.risks:
        rows.append([SEVERITY_KO[f.severity], f.title, f.category, f.quote, f.page, f.extracted_value or "", f.threshold or "", f.why, f.negotiation_point, f.basis or "",
                     {"rule": "규칙", "rule+llm": "규칙+LLM", "llm": "LLM"}[f.source]])
        fills.append(FILL[f.severity])
    _sheet(wb, "위험 조항", ["심각도", "위험", "구분", "조항 원문", "쪽", "추출값", "기준값", "왜 위험한가", "협상 포인트", "근거", "판정"], rows,
           [7, 24, 10, 60, 5, 18, 18, 44, 40, 40, 10], fills)

    rows = []
    if report.toc:
        for node in report.toc.toc:
            pts = sum(c.points for c in report.toc.scoring if c.id in node.criterion_ids and (not c.parent or c.parent not in node.criterion_ids))
            rows.append([node.number, node.title, pts, node.suggested_pages, ", ".join(node.requirement_ids), node.writing_guide])
    _sheet(wb, "배점-목차", ["번호", "장 제목", "배점", "제안 분량(쪽)", "대응 요구사항", "작성 가이드"], rows, [7, 30, 7, 12, 50, 70])

    if report.coverage:
        by_id = {r.id: r for r in report.requirements}
        rows, fills = [], []
        for it in sorted(report.coverage.items, key=lambda x: ({"missing": 0, "partial": 1, "covered": 2}[x.status], x.requirement_id)):
            r = by_id.get(it.requirement_id)
            rows.append([it.requirement_id, r.title if r else "", OBLIGATION_KO[r.obligation] if r else "", STATUS_KO[it.status], ", ".join(it.evidence_paragraph_ids),
                         it.evidence_quote or "", it.gap_note or ""])
            fills.append(FILL[it.status])
        _sheet(wb, "커버리지", ["ID", "요구사항", "필수", "판정", "근거 문단", "근거 발췌", "보완 메모"], rows, [11, 30, 6, 7, 14, 40, 50], fills)
    wb.save(path)
    return path


def write_toc_docx(report: ReviewReport, path: Path) -> Path:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Pretendard"
    style.font.size = Pt(10)
    d.add_heading(f"제안서 목차 초안 — {report.summary.title or ''}", level=0)
    if report.toc:
        d.add_paragraph(f"배점표 기준 총 {report.toc.total_pages}쪽 배분. 배점 항목 {len([c for c in report.toc.scoring if not c.parent])}개, 요구사항 {report.summary.requirement_count}건 중 미연결 {len(report.toc.unmapped_requirement_ids)}건.")
        for node in report.toc.toc:
            level = 1 if "." not in node.number else 2
            d.add_heading(f"{node.number}. {node.title} ({node.suggested_pages}쪽)", level=level)
            d.add_paragraph(node.writing_guide)
            if node.requirement_ids:
                d.add_paragraph("대응 요구사항: " + ", ".join(node.requirement_ids))
    else:
        d.add_paragraph("배점표를 찾지 못했다. toc 단계를 확인할 것.")
    d.add_heading("위험 조항 요약", level=1)
    for f in report.risks:
        p = d.add_paragraph(style="List Bullet")
        p.add_run(f"[{SEVERITY_KO[f.severity]}] {f.title} (p.{f.page}) ").bold = True
        p.add_run(f"— {f.why} 협상 포인트: {f.negotiation_point}")
    d.save(path)
    return path


def write_dashboard(report: ReviewReport, path: Path, pages: dict[int, str] | None = None) -> Path:
    """오프라인에서 열리는 단일 HTML. app.html에 report JSON을 심는다."""
    tpl = (WEB_DIR / "app.html").read_text(encoding="utf-8")
    payload = {"report": json.loads(report.model_dump_json()), "pages": pages or {}}
    data = json.dumps(payload, ensure_ascii=False).replace("</", "<\\/")
    html = tpl.replace("/*__EMBEDDED_REPORT__*/null", data)
    path.write_text(html, encoding="utf-8")
    return path
